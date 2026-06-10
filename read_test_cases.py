#!/usr/bin/env python
from openpyxl import load_workbook

# Read the main test file
wb = load_workbook('sku_files/B110_Cluster_Flashing.xlsx')
sheet = wb.active

# Get headers
headers = [str(cell.value).strip() if cell.value else '' for cell in sheet[1]]
print(f"Headers: {headers}")

# Find Test Sequence column
if 'Test Sequence' not in headers:
    print("ERROR: 'Test Sequence' column not found")
else:
    test_seq_col = headers.index('Test Sequence')
    print(f"\n'Test Sequence' is at column index: {test_seq_col}")
    print("\nTest Cases from Excel:")
    print("=" * 50)
    
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, max_row=100, values_only=True), start=2):
        test_name = row[test_seq_col] if test_seq_col < len(row) else None
        
        if test_name is None or str(test_name).strip() == "":
            continue
        
        clean_name = str(test_name).strip().replace(" ", "_")
        print(f"Row {row_idx}: '{test_name}' → Function name: '{clean_name}'")

# Also check SKU File Mapping
print("\n\n" + "=" * 50)
print("SKU File Mapping Excel Check")
print("=" * 50)

try:
    wb2 = load_workbook('SKU_File_Mapping.xlsx')
    sheet2 = wb2.active
    headers2 = [str(cell.value).strip() if cell.value else '' for cell in sheet2[1]]
    print(f"Headers in SKU Mapping: {headers2}")
except Exception as e:
    print(f"Could not read SKU_File_Mapping.xlsx: {e}")
